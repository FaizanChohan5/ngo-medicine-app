# -*- coding: utf-8 -*-

from flask import (
    Blueprint,
    render_template,
    redirect,
    url_for,
    flash,
    request,
    send_file
)

from flask_login import login_required, current_user

from datetime import datetime
from calendar import monthrange
from io import BytesIO

from app.models import (
    db,
    User,
    Medicine,
    Stock,
    StockTransaction,
    StockDispatch,
    Sale,
    SaleItem,
    Shop,
    FieldPerson
)


# ============================================================
# ADMIN BLUEPRINT
# ============================================================

admin = Blueprint(
    "admin",
    __name__,
    url_prefix="/admin"
)


# ============================================================
# ADMIN ACCESS CHECK
# ============================================================

def admin_required():

    return (
        current_user.is_authenticated
        and current_user.role == "admin"
    )


# ============================================================
# ADMIN DASHBOARD
# ============================================================

@admin.route("/dashboard")
@login_required
def dashboard():

    if not admin_required():
        return "Access Denied", 403

    return render_template(
        "admin/dashboard.html"
    )


# ============================================================
# REPORTS PAGE
# ============================================================

@admin.route("/reports")
@login_required
def reports():

    if not admin_required():
        return "Access Denied", 403

    medicines = Medicine.query.order_by(
        Medicine.id
    ).all()

    stocks = Stock.query.order_by(
        Stock.id
    ).all()

    sales = Sale.query.order_by(
        Sale.id.desc()
    ).all()

    shops = Shop.query.order_by(
        Shop.id
    ).all()

    field_persons = FieldPerson.query.order_by(
        FieldPerson.id
    ).all()

    transactions = StockTransaction.query.order_by(
        StockTransaction.id.desc()
    ).all()

    return render_template(
        "admin/reports.html",
        medicines=medicines,
        stocks=stocks,
        sales=sales,
        shops=shops,
        field_persons=field_persons,
        transactions=transactions
    )


# ============================================================
# FIELD PERSON LIST
# ============================================================

@admin.route("/field-persons")
@login_required
def field_persons():

    if not admin_required():
        return "Access Denied", 403

    persons = FieldPerson.query.order_by(
        FieldPerson.id.asc()
    ).all()

    return render_template(
        "admin/field_persons.html",
        persons=persons
    )


# ============================================================
# ADD FIELD PERSON
# ============================================================

@admin.route(
    "/field-persons/add",
    methods=["GET", "POST"]
)
@login_required
def add_field_person():

    if not admin_required():
        return "Access Denied", 403

    if request.method == "POST":

        full_name = request.form.get(
            "full_name",
            ""
        ).strip()

        username = request.form.get(
            "username",
            ""
        ).strip()

        password = request.form.get(
            "password",
            ""
        ).strip()

        phone = request.form.get(
            "phone",
            ""
        ).strip()

        if not full_name:

            flash(
                "Full name is required."
            )

            return redirect(
                url_for(
                    "admin.add_field_person"
                )
            )

        if not username:

            flash(
                "Username is required."
            )

            return redirect(
                url_for(
                    "admin.add_field_person"
                )
            )

        if not password:

            flash(
                "Password is required."
            )

            return redirect(
                url_for(
                    "admin.add_field_person"
                )
            )

        existing_user = User.query.filter_by(
            username=username
        ).first()

        if existing_user:

            flash(
                "Username already exists. "
                "Please choose another username."
            )

            return redirect(
                url_for(
                    "admin.add_field_person"
                )
            )

        try:

            user = User(
                username=username,
                role="field",
                full_name=full_name,
                is_active=True
            )

            user.set_password(password)

            db.session.add(user)

            db.session.flush()

            field_person = FieldPerson(
                user_id=user.id,
                phone=phone
            )

            db.session.add(
                field_person
            )

            db.session.commit()

            flash(
                f"Field Person '{full_name}' "
                "created successfully."
            )

            return redirect(
                url_for(
                    "admin.field_persons"
                )
            )

        except Exception as e:

            db.session.rollback()

            print(
                "FIELD PERSON CREATION ERROR:",
                e
            )

            flash(
                "Error creating Field Person. "
                "Please check the server terminal."
            )

            return redirect(
                url_for(
                    "admin.add_field_person"
                )
            )

    return render_template(
        "admin/add_field_person.html"
    )


# ============================================================
# ALL SHOPS
# ============================================================

@admin.route("/shops")
@login_required
def shops():

    if not admin_required():
        return "Access Denied", 403

    search = request.args.get(
        "search",
        ""
    ).strip()

    query = Shop.query

    if search:

        query = query.filter(
            Shop.name.ilike(
                f"%{search}%"
            )
        )

    shops = query.order_by(
        Shop.id.desc()
    ).all()

    return render_template(
        "admin/shops.html",
        shops=shops,
        search=search
    )


# ============================================================
# SHOP DETAILS
# ============================================================

@admin.route("/shops/<int:shop_id>")
@login_required
def shop_details(shop_id):

    if not admin_required():
        return "Access Denied", 403

    shop = db.session.get(
        Shop,
        shop_id
    )

    if not shop:
        return "Shop not found", 404

    sales = Sale.query.filter_by(
        shop_id=shop.id
    ).order_by(
        Sale.id.desc()
    ).all()

    return render_template(
        "admin/shop_details.html",
        shop=shop,
        sales=sales
    )


# ============================================================
# ALL SALES
# ============================================================

@admin.route("/sales")
@login_required
def sales():

    if not admin_required():
        return "Access Denied", 403

    sales = Sale.query.order_by(
        Sale.id.desc()
    ).all()

    return render_template(
        "admin/sales.html",
        sales=sales
    )


# ============================================================
# SALE DETAILS
# ============================================================

@admin.route("/sales/<int:sale_id>")
@login_required
def sale_details(sale_id):

    if not admin_required():
        return "Access Denied", 403

    sale = db.session.get(
        Sale,
        sale_id
    )

    if not sale:
        return "Sale not found", 404

    sale_items = SaleItem.query.filter_by(
        sale_id=sale.id
    ).all()

    return render_template(
        "admin/sale_details.html",
        sale=sale,
        sale_items=sale_items
    )


# ============================================================
# MONTHLY REPORT EXPORT
# ============================================================

@admin.route("/reports/export")
@login_required
def export_reports():

    if not admin_required():
        return "Access Denied", 403

    try:

        from openpyxl import Workbook

        from openpyxl.styles import (
            Font,
            Alignment,
            PatternFill,
            Border,
            Side
        )

        from openpyxl.utils import get_column_letter

    except ImportError:

        flash(
            "openpyxl is not installed in the active "
            "virtual environment."
        )

        return redirect(
            url_for("admin.reports")
        )

    # ========================================================
    # SELECT MONTH AND YEAR
    # ========================================================

    month = request.args.get(
        "month",
        datetime.now().month,
        type=int
    )

    year = request.args.get(
        "year",
        datetime.now().year,
        type=int
    )

    if month < 1 or month > 12:
        month = datetime.now().month

    if year < 2000 or year > 2100:
        year = datetime.now().year

    # ========================================================
    # DATE RANGE
    # ========================================================

    start_date = datetime(
        year,
        month,
        1
    )

    last_day = monthrange(
        year,
        month
    )[1]

    end_date = datetime(
        year,
        month,
        last_day,
        23,
        59,
        59
    )

    # ========================================================
    # DATABASE DATA
    # ========================================================

    medicines = Medicine.query.order_by(
        Medicine.id
    ).all()

    shops = Shop.query.order_by(
        Shop.id
    ).all()

    field_persons = FieldPerson.query.order_by(
        FieldPerson.id
    ).all()

    stocks = Stock.query.order_by(
        Stock.id
    ).all()

    sales = Sale.query.filter(
        Sale.created_at >= start_date,
        Sale.created_at <= end_date
    ).order_by(
        Sale.created_at
    ).all()

    sale_items = SaleItem.query.join(
        Sale,
        SaleItem.sale_id == Sale.id
    ).filter(
        Sale.created_at >= start_date,
        Sale.created_at <= end_date
    ).order_by(
        SaleItem.id
    ).all()

    dispatches = StockDispatch.query.filter(
        StockDispatch.dispatched_at >= start_date,
        StockDispatch.dispatched_at <= end_date
    ).order_by(
        StockDispatch.dispatched_at
    ).all()

    transactions = StockTransaction.query.filter(
        StockTransaction.created_at >= start_date,
        StockTransaction.created_at <= end_date
    ).order_by(
        StockTransaction.created_at
    ).all()

    # ========================================================
    # CALCULATIONS
    # ========================================================

    total_sales = sum(
        float(sale.total_amount or 0)
        for sale in sales
    )

    total_sold_quantity = sum(
        int(item.quantity or 0)
        for item in sale_items
    )

    total_dispatched_quantity = sum(
        int(dispatch.quantity or 0)
        for dispatch in dispatches
    )

    total_stock_quantity = sum(
        int(stock.quantity or 0)
        for stock in stocks
    )

    total_transactions = len(
        transactions
    )

    # ========================================================
    # CREATE WORKBOOK
    # ========================================================

    workbook = Workbook()

    default_sheet = workbook.active

    workbook.remove(
        default_sheet
    )

    # ========================================================
    # STYLES
    # ========================================================

    title_font = Font(
        bold=True,
        size=18
    )

    section_font = Font(
        bold=True,
        size=14
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    thin_side = Side(
        style="thin",
        color="B7B7B7"
    )

    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )

    center = Alignment(
        horizontal="center",
        vertical="center"
    )

    # ========================================================
    # STYLE HEADER
    # ========================================================

    def style_header(sheet, row=1):

        for cell in sheet[row]:

            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = thin_border

    # ========================================================
    # APPLY BORDERS
    # ========================================================

    def apply_borders(sheet):

        for row_cells in sheet.iter_rows():

            for cell in row_cells:

                cell.border = thin_border

    # ========================================================
    # AUTO WIDTH
    # ========================================================

    def auto_width(sheet):

        for column_cells in sheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                try:

                    value_length = len(
                        str(cell.value)
                    )

                    if value_length > max_length:
                        max_length = value_length

                except Exception:

                    pass

            sheet.column_dimensions[
                column_letter
            ].width = min(
                max(max_length + 2, 12),
                35
            )

    # ========================================================
    # SHEET 1 - MONTHLY SUMMARY
    # ========================================================

    summary = workbook.create_sheet(
        "Monthly Summary"
    )

    summary["A1"] = (
        "NGO MEDICINE MANAGEMENT SYSTEM"
    )

    summary["A1"].font = title_font

    summary["A2"] = (
        "MONTHLY OPERATIONS REPORT"
    )

    summary["A2"].font = section_font

    summary["A4"] = "Report Month"

    summary["B4"] = datetime(
        year,
        month,
        1
    ).strftime(
        "%B %Y"
    )

    summary["A5"] = "Report Period"

    summary["B5"] = (
        start_date.strftime("%d-%b-%Y")
        + " to "
        + end_date.strftime("%d-%b-%Y")
    )

    summary["A6"] = "Generated On"

    summary["B6"] = datetime.now().strftime(
        "%d-%b-%Y %H:%M"
    )

    summary["A8"] = "MONTHLY KPIs"

    summary["A8"].font = section_font

    kpis = [

        (
            "Total Medicines",
            len(medicines)
        ),

        (
            "Total Shops",
            len(shops)
        ),

        (
            "Total Field Persons",
            len(field_persons)
        ),

        (
            "Sales Transactions",
            len(sales)
        ),

        (
            "Total Medicines Sold",
            total_sold_quantity
        ),

        (
            "Total Sales Amount",
            f"Rs {total_sales:,.2f}"
        ),

        (
            "Dispatch Transactions",
            len(dispatches)
        ),

        (
            "Total Dispatched Quantity",
            total_dispatched_quantity
        ),

        (
            "Current Total Stock",
            total_stock_quantity
        ),

        (
            "Stock Transactions",
            total_transactions
        )
    ]

    row = 9

    for label, value in kpis:

        summary.cell(
            row=row,
            column=1,
            value=label
        )

        summary.cell(
            row=row,
            column=2,
            value=value
        )

        summary.cell(
            row=row,
            column=1
        ).font = Font(
            bold=True
        )

        summary.cell(
            row=row,
            column=1
        ).border = thin_border

        summary.cell(
            row=row,
            column=2
        ).border = thin_border

        row += 1

    auto_width(summary)

    # ========================================================
    # SHEET 2 - STOCK SUMMARY
    # ========================================================

    stock_sheet = workbook.create_sheet(
        "Stock Summary"
    )

    stock_sheet.append([
        "ID",
        "Medicine",
        "Category",
        "Unit",
        "Price",
        "Current Stock",
        "Status"
    ])

    style_header(
        stock_sheet
    )

    stock_map = {
        stock.medicine_id: stock.quantity
        for stock in stocks
    }

    for medicine in medicines:

        stock_sheet.append([
            medicine.id,
            medicine.name,
            medicine.category or "-",
            medicine.unit,
            medicine.price or 0,
            stock_map.get(
                medicine.id,
                0
            ),
            (
                "Active"
                if medicine.is_active
                else "Inactive"
            )
        ])

    apply_borders(
        stock_sheet
    )

    stock_sheet.freeze_panes = "A2"

    auto_width(
        stock_sheet
    )

    # ========================================================
    # SHEET 3 - SALES SUMMARY
    # ========================================================

    sales_sheet = workbook.create_sheet(
        "Sales Summary"
    )

    sales_sheet.append([
        "Sale ID",
        "Shop",
        "Field Person",
        "Total Amount",
        "Date"
    ])

    style_header(
        sales_sheet
    )

    for sale in sales:

        field_person_name = "-"

        if sale.field_person:

            if sale.field_person.user:

                field_person_name = (
                    sale.field_person.user.full_name
                )

        sales_sheet.append([
            sale.id,

            (
                sale.shop.name
                if sale.shop
                else "-"
            ),

            field_person_name,

            float(
                sale.total_amount or 0
            ),

            (
                sale.created_at.strftime(
                    "%d-%b-%Y %H:%M"
                )
                if sale.created_at
                else "-"
            )
        ])

    apply_borders(
        sales_sheet
    )

    sales_sheet.freeze_panes = "A2"

    auto_width(
        sales_sheet
    )

    # ========================================================
    # SHEET 4 - MEDICINE SALES
    # ========================================================

    items_sheet = workbook.create_sheet(
        "Medicine Sales"
    )

    items_sheet.append([
        "ID",
        "Sale ID",
        "Shop",
        "Medicine",
        "Quantity",
        "Rate",
        "Total",
        "Date"
    ])

    style_header(
        items_sheet
    )

    for item in sale_items:

        sale = item.sale

        items_sheet.append([
            item.id,

            item.sale_id,

            (
                sale.shop.name
                if sale and sale.shop
                else "-"
            ),

            (
                item.medicine.name
                if item.medicine
                else "-"
            ),

            item.quantity,

            item.rate,

            item.total,

            (
                sale.created_at.strftime(
                    "%d-%b-%Y %H:%M"
                )
                if sale and sale.created_at
                else "-"
            )
        ])

    apply_borders(
        items_sheet
    )

    items_sheet.freeze_panes = "A2"

    auto_width(
        items_sheet
    )

    # ========================================================
    # SHEET 5 - DISPATCH SUMMARY
    # ========================================================

    dispatch_sheet = workbook.create_sheet(
        "Dispatch Summary"
    )

    dispatch_sheet.append([
        "ID",
        "Field Person",
        "Medicine",
        "Quantity",
        "Date"
    ])

    style_header(
        dispatch_sheet
    )

    for dispatch in dispatches:

        field_person_name = "-"

        if dispatch.field_person:

            if dispatch.field_person.user:

                field_person_name = (
                    dispatch.field_person.user.full_name
                )

        dispatch_sheet.append([
            dispatch.id,

            field_person_name,

            (
                dispatch.medicine.name
                if dispatch.medicine
                else "-"
            ),

            dispatch.quantity,

            (
                dispatch.dispatched_at.strftime(
                    "%d-%b-%Y %H:%M"
                )
                if dispatch.dispatched_at
                else "-"
            )
        ])

    apply_borders(
        dispatch_sheet
    )

    dispatch_sheet.freeze_panes = "A2"

    auto_width(
        dispatch_sheet
    )

    # ========================================================
    # SHEET 6 - STOCK TRANSACTIONS
    # ========================================================

    transaction_sheet = workbook.create_sheet(
        "Stock Transactions"
    )

    transaction_sheet.append([
        "ID",
        "Medicine",
        "Transaction Type",
        "Quantity",
        "Performed By",
        "Date"
    ])

    style_header(
        transaction_sheet
    )

    for transaction in transactions:

        transaction_sheet.append([
            transaction.id,

            (
                transaction.medicine.name
                if transaction.medicine
                else "-"
            ),

            transaction.transaction_type,

            transaction.quantity,

            (
                transaction.user.full_name
                if transaction.user
                else "-"
            ),

            (
                transaction.created_at.strftime(
                    "%d-%b-%Y %H:%M"
                )
                if transaction.created_at
                else "-"
            )
        ])

    apply_borders(
        transaction_sheet
    )

    transaction_sheet.freeze_panes = "A2"

    auto_width(
        transaction_sheet
    )

    # ========================================================
    # SHEET 7 - SHOPS
    # ========================================================

    shop_sheet = workbook.create_sheet(
        "Shops"
    )

    shop_sheet.append([
        "ID",
        "Shop Name",
        "Owner Name",
        "Phone",
        "Address",
        "Field Person"
    ])

    style_header(
        shop_sheet
    )

    for shop in shops:

        field_person_name = "-"

        if shop.field_person:

            if shop.field_person.user:

                field_person_name = (
                    shop.field_person.user.full_name
                )

        shop_sheet.append([
            shop.id,
            shop.name,
            shop.owner_name or "-",
            shop.phone or "-",
            shop.address or "-",
            field_person_name
        ])

    apply_borders(
        shop_sheet
    )

    shop_sheet.freeze_panes = "A2"

    auto_width(
        shop_sheet
    )

    # ========================================================
    # SAVE WORKBOOK IN MEMORY
    # ========================================================

    output = BytesIO()

    workbook.save(
        output
    )

    output.seek(0)

    filename = (
        f"NGO_Monthly_Report_"
        f"{year}_{month:02d}.xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )